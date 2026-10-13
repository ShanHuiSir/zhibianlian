"""Bounded document parsing. Original positions survive field mapping."""
import csv
import io
import re
import zipfile
from pathlib import Path
import fitz
from docx import Document
from openpyxl import load_workbook

ROLES = {'old': '旧图纸', 'new': '新图纸', 'ecn': '变更通知', 'bom': 'BOM',
         'inspection': '检验标准', 'inventory': '库存', 'wip': '在制品', 'orders': '订单', 'process': '工艺说明'}
REQUIRED = {
    'bom': ['record_id', 'product_no', 'product_rev', 'part_no', 'revision', 'quantity'],
    'inspection': ['item_id', 'part_no', 'feature_id', 'revision', 'nominal', 'unit'],
    'inventory': ['record_id', 'batch', 'part_no', 'revision', 'quantity'],
    'wip': ['record_id', 'work_order', 'part_no', 'revision', 'quantity'],
    'orders': ['record_id', 'order_id', 'part_no', 'revision', 'quantity'],
}
ALIASES = {'编号':'record_id','记录编号':'record_id','检验项':'item_id','料号':'part_no','零件编号':'part_no',
           '版本':'revision','数量':'quantity','批次':'batch','工单':'work_order','订单号':'order_id',
           '特征':'feature_id','名义值':'nominal','下偏差':'lower','上偏差':'upper','单位':'unit',
           '产品编号':'product_no','产品版本':'product_rev'}

def parse_file(path):
    path = Path(path)
    ext = path.suffix.lower()
    if path.stat().st_size > 20 * 1024 * 1024:
        raise ValueError('每份文件限 20 MB')
    pages, rows, headers = [], [], []
    if ext in ('.docx', '.xlsx'):
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024 or len(z.infolist()) > 3000:
                raise ValueError('文件解压后超限')
    if ext == '.pdf':
        with fitz.open(path) as doc:
            if doc.needs_pass or len(doc) > 30:
                raise ValueError('PDF 加密或超过 30 页')
            for p in doc:
                pages.append({'location': f'第 {p.number+1} 页', 'page': p.number+1, 'text': p.get_text()})
    elif ext in ('.txt', '.csv'):
        raw = path.read_text(encoding='utf-8-sig')
        if '\x00' in raw:
            raise ValueError('请提供 UTF-8 文本')
        if ext == '.csv':
            reader = csv.DictReader(io.StringIO(raw))
            headers = reader.fieldnames or []
            if len(set(headers)) != len(headers):
                raise ValueError('表头重复')
            for row in reader:
                if None in row:
                    raise ValueError('CSV 行列数不匹配')
                rows.append({'location': f'CSV 记录结束行 {reader.line_num}', 'values': {k: (v or '').strip() for k,v in row.items()}})
        else:
            pages = [{'location':f'段落 {i+1}', 'text': t} for i,t in enumerate(raw.splitlines()) if t.strip()]
    elif ext == '.xlsx':
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for ws in wb:
                it = ws.iter_rows()
                first = next(it, ())
                hs = [str(c.value).strip() if c.value is not None else '' for c in first]
                if not hs or not any(hs): continue
                if len(set(hs)) != len(hs) or any(not x for x in hs):
                    raise ValueError(f'{ws.title} 存在空或重复表头')
                headers.extend(h for h in hs if h not in headers)
                for i, cells in enumerate(it, 2):
                    if all(c.value is None for c in cells): continue
                    if any(c.data_type == 'f' for c in cells):
                        raise ValueError(f'{ws.title} 第 {i} 行含公式，请转为值后导入')
                    vals = {h: '' if c.value is None else str(c.value).strip() for h,c in zip(hs,cells)}
                    rows.append({'location': f'{ws.title} 第 {i} 行', 'values': vals})
                    if len(rows) > 1000: raise ValueError('表格超过 1000 条记录')
        finally: wb.close()
    elif ext == '.docx':
        d = Document(path)
        pages = [{'location':f'段落 {i+1}', 'text':p.text} for i,p in enumerate(d.paragraphs) if p.text.strip()]
        for i,t in enumerate(d.tables):
            for j,r in enumerate(t.rows):
                pages.append({'location':f'表 {i+1} 第 {j+1} 行', 'text':' | '.join(c.text for c in r.cells)})
    else:
        raise ValueError('仅支持 PDF / XLSX / CSV / DOCX / TXT')
    if len(rows)>1000: raise ValueError('表格超过 1000 条记录')
    if not pages and not rows: raise ValueError('未找到有效资料')
    return {'pages':pages,'rows':rows,'headers':headers,
            'mapping':{h: ALIASES.get(h,h) for h in headers}, 'mapping_confirmed':not rows}

def local_drawing(pages):
    """Explicit interchange notation; no geometry inference or preset results."""
    text = '\n'.join(p['text'] for p in pages)
    part = re.search(r'PART\s*[:=]\s*([^\s]+)',text)
    rev = re.search(r'REV\s*[:=]\s*([^\s]+)',text)
    features=[]
    for p in pages:
        for line in p['text'].splitlines():
            m=re.fullmatch(r'\s*FEATURE\s+(\S+)\s+(diameter|length)\s+([+\-\d.]+)\s+(mm|cm|m|inch)\s+LOW\s+([+\-\d.]+|unknown)\s+HIGH\s+([+\-\d.]+|unknown)\s*',line)
            if m:
                fid,kind,n,u,lo,hi=m.groups()
                features.append(dict(feature_id=fid,kind=kind,nominal=n,unit=u,lower=None if lo=='unknown' else lo,
                                     upper=None if hi=='unknown' else hi,page=p.get('page',1),quote=line.strip()))
    if not part or not rev or not features:
        raise ValueError('本地模式只支持样例中的标准 FEATURE 标注。普通图纸请切换 DeepSeek 识图，或补充结构化标注。')
    return {'part_no':part.group(1),'revision':rev.group(1),'features':features}
